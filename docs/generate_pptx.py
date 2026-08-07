import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(11, 15, 25)
    CARD_BG = RGBColor(18, 26, 43)
    ACCENT_INDIGO = RGBColor(99, 102, 241)
    ACCENT_CYAN = RGBColor(6, 182, 212)
    ACCENT_GREEN = RGBColor(16, 185, 129)
    ACCENT_WARNING = RGBColor(245, 158, 11)
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(156, 163, 175)

    blank_slide_layout = prs.slide_layouts[6]

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK

    def add_header(slide, title_text, badge_text="MINDSYNC AI PRESENTATION"):
        # Header Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.5), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(30, 27, 75)
        badge.line.color.rgb = ACCENT_INDIGO
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = ACCENT_CYAN
        p_b.alignment = PP_ALIGN.CENTER

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide1)

    # Main Card
    card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = ACCENT_INDIGO

    tb = slide1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "MindSync AI"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE

    p1 = tf.add_paragraph()
    p1.text = "An AI-Driven Mental Wellness Companion for IT Professionals"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "Predictive Machine Learning • Gemini LLM Orchestration • Offline-First Mobile Design"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "Presenter: Yuvin Chandula   |   Module: CMP7003 Emerging Mobile Applications"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_GREEN
    p3.space_before = Pt(35)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide2)
    add_header(slide2, "The IT Developer Burnout Crisis & Industry Problem")

    # Left Card
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_INDIGO
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "⚠️ The IT Workplace Burnout Crisis"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(239, 68, 68)

    bullets1 = [
      "High cognitive fatigue, tight sprint deadlines, and prolonged screen exposure.",
      "Over 60% of software engineers report experiencing severe burnout during their career.",
      "Lack of tools connecting cognitive work fatigue to physical lifestyle telemetry."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    # Right Card
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_INDIGO
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "⚡ Limitations of Existing Wellness Apps"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_WARNING

    bullets2 = [
      "Generic fitness apps track steps but ignore developer work stress and mood ratings.",
      "Siloed logs: Sleep, hydration, work hours, and stress are never correlated.",
      "Static push notifications cause alert fatigue and are quickly disabled.",
      "Black-box wellness advice lacks transparency and actionable explanations."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Aim & Objectives
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide3)
    add_header(slide3, "Project Aim & Strategic Technical Objectives")

    aim_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.3))
    aim_card.fill.solid()
    aim_card.fill.fore_color.rgb = CARD_BG
    aim_card.line.color.rgb = ACCENT_CYAN
    tf_a = aim_card.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "🎯 Core Project Aim"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p2 = tf_a.add_paragraph()
    p2.text = "To design, engineer, and evaluate an offline-first mobile companion that accurately predicts developer burnout risks, explains key risk drivers via SHAP values, and delivers context-aware AI coaching."
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(6)

    # 4 Objective Boxes
    objs = [
      ("1. Flutter Mobile App", "Cross-platform client with Material 3, Riverpod state management, and fl_chart trend analytics."),
      ("2. Async FastAPI Backend", "High-performance microservices engine delivering <22ms inference latencies."),
      ("3. Random Forest Classifier", "Scikit-Learn predictive model achieving 93.3% accuracy with SHAP explainability."),
      ("4. Offline Hive Resilience", "Local NoSQL storage with automatic Firestore cloud sync on connection recovery.")
    ]
    for i, (title, desc) in enumerate(objs):
        x = Inches(0.8 + (i % 2) * 6.0)
        y = Inches(3.3 + (i // 2) * 1.8)
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(1.6))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_INDIGO
        tf_o = box.text_frame
        tf_o.word_wrap = True
        p_t = tf_o.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_GREEN
        p_d = tf_o.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 4: Architecture
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide4)
    add_header(slide4, "Clean Architecture & Domain-Driven System Design")

    layers = [
      ("Presentation Layer", "Flutter Widgets, Riverpod Controllers, GoRouter Navigation", ACCENT_INDIGO),
      ("Domain & Data Layer", "Entities, Use Cases, Hive DB Local Storage, Sync Engine", ACCENT_CYAN),
      ("Backend Microservices", "FastAPI Endpoints, Random Forest ML Engine, Gemini LLM", ACCENT_GREEN)
    ]
    for i, (title, desc, color) in enumerate(layers):
        x = Inches(0.8 + i * 4.0)
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        tf_l = box.text_frame
        tf_l.word_wrap = True
        p_t = tf_l.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = color
        p_d = tf_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 5: Tech Stack Table
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide5)
    add_header(slide5, "Technology Stack & Infrastructure Specifications")

    rows = [
      ["Mobile Client", "Flutter / Dart", "^3.22 / ^3.4", "Cross-platform UI framework with Material 3 design"],
      ["State Management", "Riverpod", "^2.5.1", "Compile-safe state container & dependency injection"],
      ["Local Database", "Hive DB", "^1.1.0", "High-performance NoSQL local key-value storage"],
      ["Backend API", "FastAPI / Python", "^0.111 / ^3.11", "Async Python microservices server for ML endpoints"],
      ["Predictive ML", "Scikit-Learn & SHAP", "^1.5.0", "Random Forest Classifier with SHAP explainability"],
      ["Generative AI", "Google Gemini 1.5 Flash", "REST API", "Context-aware developer coaching & sentiment analysis"],
      ["Cloud Suite", "Firebase Platform", "^6.5.0 Admin", "Firebase Auth, Cloud Firestore, Firebase Storage, FCM"]
    ]
    table_shape = slide5.shapes.add_table(8, 4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table

    headers = ["Domain", "Technology", "Version", "Role & Functionality"]
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 27, 75)
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

    for row_idx, data in enumerate(rows, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE if col_idx < 2 else TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 6: ML Pipeline
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide6)
    add_header(slide6, "Machine Learning Pipeline & SHAP Explainability")

    # Left Card
    c1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_INDIGO
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🌲 Random Forest Prediction Engine"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    bullets1 = [
      "Model Choice: Random Forest Classifier (Scikit-Learn).",
      "Selected for superior performance and robustness on tabular health metrics.",
      "Input Vector: Sleep hours, Work/Screen time, Mood rating (1-5), Stress level (1-5), Active steps, Water intake.",
      "Model Accuracy: 93.3% accuracy on cross-validation sets."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    # Right Card
    c2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_INDIGO
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🔍 SHAP Value Explainability (XAI)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    bullets2 = [
      "Prevents 'black-box AI' by calculating SHAP tree feature attribution weights.",
      "Translates complex tree decisions into human-understandable risk drivers.",
      "Identifies top 3 factors for each prediction.",
      "Example output: 'High risk driven by <5.5 hrs sleep and >10 hrs work time'."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 7: AI Orchestration
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide7)
    add_header(slide7, "Gemini LLM Orchestration & Fail-Safe Architecture")

    box1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    box1.fill.solid()
    box1.fill.fore_color.rgb = CARD_BG
    box1.line.color.rgb = ACCENT_CYAN
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "💬 Gemini 1.5 Flash Integration"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    bullets1 = [
      "Context-Aware Habit Suggestions: Generates actionable advice based on daily metrics.",
      "Wellness AI Coach Chat: Provides supportive developer chat coaching.",
      "Sentiment Analysis: Evaluates daily text journal entries for emotional state."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    box2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    box2.fill.solid()
    box2.fill.fore_color.rgb = CARD_BG
    box2.line.color.rgb = ACCENT_GREEN
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛡️ Resilience & Safety Mechanisms"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    bullets2 = [
      "Exponential Backoff: Automatic retry loop on Gemini API rate limits or network drops.",
      "Fail-Safe Fallback Engine: Instantly serves rule-based local summaries if Gemini is offline.",
      "Prompt Sanitization: Protects against prompt injection attacks.",
      "Medical Disclaimer: Mandatory disclaimer attached to all wellness coaching advice."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 8: Mobile Features
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide8)
    add_header(slide8, "Mobile Application Features & UI/UX Design")

    feats = [
      ("📊 Dynamic Dashboard", "Real-time burnout risk gauge, daily wellness score cards, and location-aware weather widget."),
      ("📝 Interactive Logger", "Drag sliders to log sleep hours, mood rating, stress level, steps, and water intake."),
      ("📈 fl_chart Telemetry", "Interactive weekly and monthly trend graphs comparing mood against burnout risk factors."),
      ("📄 PDF & JSON Exporter", "One-click export of complete encrypted wellness records for personal backup.")
    ]
    for i, (title, desc) in enumerate(feats):
        x = Inches(0.8 + (i % 2) * 6.0)
        y = Inches(1.8 + (i // 2) * 2.5)
        box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_INDIGO
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 9: Offline First
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide9)
    add_header(slide9, "Offline-First Architecture & Dynamic Sync Engine")

    box1 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    box1.fill.solid()
    box1.fill.fore_color.rgb = CARD_BG
    box1.line.color.rgb = ACCENT_INDIGO
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "💾 Synchronous Hive Local DB"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    bullets1 = [
      "Synchronous Write: Every user entry commits immediately to high-speed Hive key-value boxes.",
      "Full Usability: 100% feature functionality available without an active internet connection.",
      "Instant UI Response: Eliminates network loading spinners during logging."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    box2 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    box2.fill.solid()
    box2.fill.fore_color.rgb = CARD_BG
    box2.line.color.rgb = ACCENT_GREEN
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🔄 Background Sync Engine"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    bullets2 = [
      "Connectivity Observer: Monitors network state changes in real time using connectivity_plus.",
      "Sync Queue Manager: Maintains local queue of offline log entries with timestamps.",
      "Automatic Sync: Batch uploads pending logs to Cloud Firestore when internet restores.",
      "Verified Result: 0% data loss across all offline simulation test scenarios."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 10: Security & Quality
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide10)
    add_header(slide10, "Security Hardening & Quality Assurance")

    sec_items = [
      ("🔐 Authentication & Data Isolation", "Firebase JWT ID Token verification in FastAPI backend middleware. Firestore Security Rules enforce strict row-level user data isolation."),
      ("🛡️ OWASP API Protection", "OWASP secure headers (HSTS, CSP, X-Frame-Options), memory rate-limiting middleware, and Pydantic schema validation."),
      ("🧪 Comprehensive Automated Testing", "Automated Pytest suite verifying backend endpoints and ML models. Flutter Test widget and Riverpod provider testing.")
    ]
    for i, (title, desc) in enumerate(sec_items):
        box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i * 1.6), Inches(11.7), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_CYAN
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 11: Performance Evaluation
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide11)
    add_header(slide11, "Quantitative System Performance Evaluation")

    metrics = [
      ("93.3%", "ML Model Accuracy"),
      ("<22 ms", "API Inference Latency"),
      ("1.1 s", "Mobile App Cold Boot"),
      ("60 FPS", "UI Render Frame Rate")
    ]
    for i, (val, label) in enumerate(metrics):
        x = Inches(0.8 + i * 3.0)
        box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.7), Inches(2.0))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_GREEN
        tf = box.text_frame
        tf.word_wrap = True
        p_v = tf.paragraphs[0]
        p_v.text = val
        p_v.font.size = Pt(36)
        p_v.font.bold = True
        p_v.font.color.rgb = ACCENT_GREEN
        p_v.alignment = PP_ALIGN.CENTER
        p_l = tf.add_paragraph()
        p_l.text = label
        p_l.font.size = Pt(12)
        p_l.font.color.rgb = TEXT_MUTED
        p_l.alignment = PP_ALIGN.CENTER
        p_l.space_before = Pt(8)

    summary_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.4))
    summary_box.fill.solid()
    summary_box.fill.fore_color.rgb = CARD_BG
    summary_box.line.color.rgb = ACCENT_INDIGO
    tf_s = summary_box.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "📊 Key Benchmark Insights"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    bullets = [
      "FastAPI Async Microservices maintain sub-22ms prediction latencies under load.",
      "Hive local DB enables instant startup times (1.1s) and zero frame drops during logging.",
      "100% sync fidelity verified during network disconnect/reconnect simulation tests."
    ]
    for b in bullets:
        p = tf_s.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 12: Viva Script
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide12)
    add_header(slide12, "10-Minute System Viva Demonstration Guide")

    steps = [
      ("Min 0-2: Auth & Onboarding", "Launch app -> Splash screen -> Firebase Auth registration/login -> Redirect to dashboard."),
      ("Min 2-4: Dashboard Telemetry", "Explain wellness score cards, burnout risk gauge, and location weather widget."),
      ("Min 4-6: Mood Log & ML Prediction", "Drag sliders to log daily metrics -> Save entry -> Watch Random Forest update risk + SHAP explanation."),
      ("Min 6-8: AI Coach & Recommendations", "View SHAP recommendations -> Chat with Gemini AI Coach -> Verify medical disclaimers."),
      ("Min 8-9: Offline Mode Test", "Disable Wi-Fi -> Log entry in Airplane mode -> Verify local Hive write -> Re-enable Wi-Fi -> Observe cloud sync."),
      ("Min 9-10: Analytics & Export", "Navigate to Reports -> View fl_chart trend lines -> Go to Settings -> Export PDF report backup.")
    ]
    for i, (title, desc) in enumerate(steps):
        x = Inches(0.8 + (i % 2) * 6.0)
        y = Inches(1.8 + (i // 2) * 1.7)
        box = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(1.5))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_INDIGO
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 13: Roadmap
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide13)
    add_header(slide13, "Future Horizons & Academic Research Roadmap")

    cards = [
      ("⌚ Wearables Integration", "Direct integration with Apple HealthKit and Google Health Connect for automated step, heart rate, and sleep sync."),
      ("🔒 Federated Learning", "On-device federated learning allowing predictive models to train across users without exposing raw health telemetry."),
      ("🌐 Multilingual AI Coaching", "Expanding Gemini LLM prompt templates to support multilingual developer teams across global engineering hubs.")
    ]
    for i, (title, desc) in enumerate(cards):
        x = Inches(0.8 + i * 4.0)
        box = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_INDIGO
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 14: Conclusion
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide14)
    add_header(slide14, "Conclusion & Q&A Board Discussion")

    c_box = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.8))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_BG
    c_box.line.color.rgb = ACCENT_GREEN
    tf_c = c_box.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "Summary of Achievements"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    achievements = [
      "Engineered a production-grade, offline-first mental health platform for IT professionals.",
      "Combined 93.3% accurate Random Forest predictions with transparent SHAP explanations.",
      "Delivered fail-safe Gemini LLM orchestration with automatic local fallback engines.",
      "Demonstrated 100% sync reliability with zero data loss in offline testing."
    ]
    for a in achievements:
        p = tf_c.add_paragraph()
        p.text = "✓ " + a
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    p_end = tf_c.add_paragraph()
    p_end.text = "Thank You! Open for Questions & Board Discussion."
    p_end.font.size = Pt(18)
    p_end.font.bold = True
    p_end.font.color.rgb = ACCENT_CYAN
    p_end.space_before = Pt(24)

    output_path = "docs/MindSync_AI_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
