"""
MindSync AI - User Guide Document Generator with Embedded Screenshots
Generates a Word document (.docx) with attached real screenshots for each feature:
  - Cover Page & Introduction
  - Getting Started & Login
  - Dashboard & Weather Integration
  - Mood Journaling & Telemetry
  - AI Wellness Coach (Gemini Chat)
  - Reports & Analytics (Charts & Export)
  - Profile Management
  - Settings & Security
  - Auto Sign-Out on Inactivity
  - Offline Mode & Synchronization
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
import os

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_styled_table(doc, headers, rows, header_color="1E88E5"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        set_cell_shading(hdr_cells[i], header_color)

    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_text in enumerate(row_data):
            row_cells[col_idx].text = str(cell_text)
            for paragraph in row_cells[col_idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        if row_idx % 2 == 1:
            for cell in row_cells:
                set_cell_shading(cell, "F0F4F8")

    return table

def add_screenshot_figure(doc, img_path, caption_text, width_inches=3.2):
    """Adds a screenshot image centered with a figure caption."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run(f"Figure: {caption_text}")
        run_cap.italic = True
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(100, 100, 100)
        doc.add_paragraph()
    else:
        p_err = doc.add_paragraph()
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_err = p_err.add_run(f"[Screenshot File Not Found: {img_path}]")
        run_err.italic = True
        run_err.font.color.rgb = RGBColor(200, 0, 0)

def create_user_guide():
    doc = Document()
    artifacts_dir = r"C:\Users\ychandula\.gemini\antigravity-ide\brain\a2180a62-f1ed-4e1e-8b8f-0124945d4f1e"
    local_img_dir = os.path.join(os.path.dirname(__file__), "user_guide_screenshots")

    # Margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    for _ in range(5):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MindSync AI")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(30, 136, 229)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Official User Guide & Visual Feature Walkthrough")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("━" * 50)
    run.font.color.rgb = RGBColor(30, 136, 229)

    doc.add_paragraph()

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = desc.add_run("A Step-by-Step Practical Manual with Real Mobile App Screenshots")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(80, 80, 80)

    for _ in range(4):
        doc.add_paragraph()

    info_items = [
        ("Application:", "MindSync AI Mobile Client"),
        ("Target Platform:", "Android & iOS (Pixel 7 Tested)"),
        ("Author:", "Yuvin Chandula"),
        ("Date:", "July 2026"),
    ]
    for label, value in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label + " ")
        run.bold = True
        run.font.size = Pt(11)
        run = p.add_run(value)
        run.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # TABLE OF CONTENTS
    # =========================================================================
    doc.add_heading('Table of Contents', level=1)
    toc_items = [
        "1. Welcome to MindSync AI",
        "2. Installation & Setup",
        "3. User Authentication & Security",
        "4. Navigation & Dashboard",
        "5. Mood Journaling & Wellness Telemetry",
        "6. AI Wellness Coach (Gemini Chat)",
        "7. Reports & Health Analytics",
        "8. Profile & Account Settings",
        "9. Auto Sign-Out & Security Features",
        "10. Offline Mode & Synchronization",
        "11. Summary & User Support",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        for run in p.runs:
            run.bold = True

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 1: WELCOME
    # =========================================================================
    doc.add_heading('1. Welcome to MindSync AI', level=1)
    doc.add_paragraph(
        'Welcome to MindSync AI! MindSync AI is an intelligent mental wellness companion application '
        'crafted specifically for software developers, IT engineers, system admins, and tech workers. '
        'Using machine learning predictions and AI coaching, MindSync AI helps you manage workplace stress, '
        'track daily wellness habits, prevent burnout, and receive context-aware wellness nudges.'
    )
    doc.add_paragraph(
        'This User Guide will walk you through every screen and feature of the application with clear '
        'instructions and attached mobile screenshots.'
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 2: SETUP
    # =========================================================================
    doc.add_heading('2. Installation & Setup', level=1)
    doc.add_paragraph(
        'MindSync AI is available for Android and iOS devices. Follow these quick steps to get started:'
    )
    steps = [
        'Launch the MindSync AI app on your mobile device or emulator.',
        'If launching for the first time, review the 3-slide onboarding carousel introducing Mood Tracking, Burnout Prediction, and AI Coaching.',
        'Tap "Get Started" to open the Login & Registration screen.',
    ]
    for i, s in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {s}')

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 3: AUTHENTICATION
    # =========================================================================
    doc.add_heading('3. User Authentication & Security', level=1)
    doc.add_paragraph(
        'MindSync AI provides secure authentication backed by Firebase Auth and custom JWT tokens.'
    )
    doc.add_heading('3.1 Signing In', level=2)
    doc.add_paragraph(
        '1. Enter your registered email address (e.g. yuvinchandula4@gmail.com).\n'
        '2. Enter your password in the Password field (tap the eye icon to toggle visibility).\n'
        '3. Tap the blue "Sign In" button to access your account.'
    )

    login_img = os.path.join(local_img_dir, "01_login_screen.png")
    if not os.path.exists(login_img):
        login_img = os.path.join(artifacts_dir, "auth_test_screen_2.png")
    add_screenshot_figure(doc, login_img, "MindSync AI Login & Authentication Screen")

    doc.add_heading('3.2 Creating an Account', level=2)
    doc.add_paragraph(
        'If you do not have an account, tap "Register" at the bottom of the sign-in screen. '
        'Fill in your Full Name, Email, Occupation, and Password to create your profile.'
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 4: DASHBOARD
    # =========================================================================
    doc.add_heading('4. Navigation & Dashboard', level=1)
    doc.add_paragraph(
        'The Dashboard is your main wellness hub. It combines live weather metrics, overall wellness score, '
        'burnout risk calculations, and quick access buttons.'
    )

    dash_img = os.path.join(artifacts_dir, "full_feature_test_2.png")
    if not os.path.exists(dash_img):
        dash_img = os.path.join(artifacts_dir, "weather_final_success.png")
    add_screenshot_figure(doc, dash_img, "MindSync AI Main Dashboard Hub with Live Weather & Risk Index")

    doc.add_heading('4.1 Dashboard Components', level=2)
    components = [
        ('Header & Profile Avatar', 'Displays your name, date, notification bell (with unread count badge), and settings icon.'),
        ('Weather Widget', 'Fetches live regional weather (temperature, humidity, condition) and gives a contextual nudge like "Overcast weather. Great for a screen break near a window!".'),
        ('Today\'s Wellness Score', 'Calculates a 0-100 score based on your sleep, steps, and activity levels (e.g. 82/100 - Excellent).'),
        ('Burnout Risk Index', 'Displays your current ML-calculated burnout risk percentage (e.g. 42% Moderate Risk) with a "Details" button for deep SHAP factor analysis.'),
        ('Bottom Navigation Bar', 'Provides 5 quick-switch tabs: Dashboard, Mood, AI Chat, Reports, and Profile.'),
    ]
    for title, desc in components:
        p = doc.add_paragraph()
        run = p.add_run(f'• {title}: ')
        run.bold = True
        p.add_run(desc)

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 5: MOOD JOURNALING
    # =========================================================================
    doc.add_heading('5. Mood Journaling & Wellness Telemetry', level=1)
    doc.add_paragraph(
        'Logging your daily mood and stress metrics allows MindSync AI to predict burnout risk accurately.'
    )

    mood_img = os.path.join(artifacts_dir, "mood_screen.png")
    add_screenshot_figure(doc, mood_img, "Mood Journaling Screen with Stress Sliders & Tags")

    doc.add_heading('5.1 How to Log a Mood Entry', level=2)
    mood_steps = [
        'Tap the "Mood" tab on the bottom navigation bar.',
        'Select your mood emoji (Very Happy, Happy, Neutral, Sad, Anxious, Exhausted, Angry).',
        'Adjust the Stress Level slider from 1 (Calm) to 10 (Severe).',
        'Adjust the Energy Level slider from 1 (Exhausted) to 10 (Vibrant).',
        'Type optional journal notes describing your day or work pressure.',
        'Select context tags like "work", "meeting", "deadline", or "exercise".',
        'Tap "Save Entry" to update your metrics and recalculate burnout risk.',
    ]
    for i, s in enumerate(mood_steps, 1):
        doc.add_paragraph(f'{i}. {s}')

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 6: AI CHAT
    # =========================================================================
    doc.add_heading('6. AI Wellness Coach (Gemini Chat)', level=1)
    doc.add_paragraph(
        'The AI Chat feature connects you with a personal wellness coach powered by Google Gemini 1.5 Flash. '
        'It understands IT professional stressors and provides practical, empathetic advice.'
    )

    chat_img = os.path.join(artifacts_dir, "chat_final_verified.png")
    if not os.path.exists(chat_img):
        chat_img = os.path.join(artifacts_dir, "ai_chat_screen.png")
    add_screenshot_figure(doc, chat_img, "AI Wellness Coach Conversation powered by Google Gemini")

    doc.add_heading('6.1 Chat Features', level=2)
    chat_features = [
        ('Interactive AI Dialogue', 'Ask questions about screen fatigue, posture, stress management, or sleep habits.'),
        ('Prompt Chips', 'Tap quick suggestions like "Tips for screen fatigue", "Help me wind down", or "Breathing exercises".'),
        ('Markdown Formatting', 'AI responses include bullet points, bold key takeaways, and clear instructions.'),
        ('Clear Chat', 'Tap the trash icon in the header to reset conversation history.'),
    ]
    for title, desc in chat_features:
        p = doc.add_paragraph()
        run = p.add_run(f'• {title}: ')
        run.bold = True
        p.add_run(desc)

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 7: REPORTS
    # =========================================================================
    doc.add_heading('7. Reports & Health Analytics', level=1)
    doc.add_paragraph(
        'The Reports screen provides visual trend graphs and summary analytics of your mental wellness over time.'
    )

    reports_img = os.path.join(artifacts_dir, "reports_screen.png")
    add_screenshot_figure(doc, reports_img, "Interactive Reports & Analytics Dashboard")

    doc.add_heading('7.1 Key Report Features', level=2)
    report_items = [
        'Filter trends by 7 Days, 30 Days, or 90 Days.',
        'Interactive mood & stress line graphs comparing daily ratings.',
        'Sleep & step bar charts tracking active movement and rest habits.',
        'Export report data as clean PDF or CSV files using the export button.',
    ]
    for f in report_items:
        doc.add_paragraph(f, style='List Bullet')

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 8: PROFILE
    # =========================================================================
    doc.add_heading('8. Profile & Account Settings', level=1)
    doc.add_paragraph(
        'Manage your personal information, job title, and settings from the Profile screen.'
    )

    profile_img = os.path.join(artifacts_dir, "profile_screen.png")
    add_screenshot_figure(doc, profile_img, "My Profile Management Screen")

    doc.add_heading('8.1 Editing Profile Details', level=2)
    doc.add_paragraph(
        'Tap the "Edit" button in the upper right corner to modify your Full Name, Occupation '
        '(e.g., System Engineer), Bio, Phone Number, Date of Birth, Gender, and Country. '
        'You can also tap the camera icon on the avatar to upload a custom profile picture.'
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 9: AUTO SIGN-OUT
    # =========================================================================
    doc.add_heading('9. Auto Sign-Out & Security Features', level=1)
    doc.add_paragraph(
        'MindSync AI includes advanced security controls to protect your private health data on unattended devices.'
    )

    security_img = os.path.join(artifacts_dir, "auto_signout_switch_snackbar.png")
    if not os.path.exists(security_img):
        security_img = os.path.join(artifacts_dir, "security_settings_final_screen.png")
    add_screenshot_figure(doc, security_img, "Security Settings with Auto Sign-Out Controls & Activity Selector")

    doc.add_heading('9.1 Auto Sign-Out Configuration', level=2)
    doc.add_paragraph(
        '1. Open Profile -> Tap Settings Icon -> Select "Security & Password".\n'
        '2. Toggle "Auto Sign-Out on Inactivity" ON.\n'
        '3. Select an inactivity timeout duration: 1 min (Demo Mode), 5 mins, 15 mins, 30 mins, or 60 mins.\n'
        '4. If no touch or scroll activity occurs during this period, the app automatically logs out and displays an orange notification snackbar.'
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 10: OFFLINE MODE
    # =========================================================================
    doc.add_heading('10. Offline Mode & Synchronization', level=1)
    doc.add_paragraph(
        'MindSync AI is built with an offline-first architecture using local Hive storage. '
        'If network connection drops, the app seamlessly switches to cached metrics.'
    )

    offline_img = os.path.join(artifacts_dir, "offline_screen.png")
    add_screenshot_figure(doc, offline_img, "Offline Fallback Mode displaying cached metrics")

    doc.add_heading('10.1 How Offline Sync Works', level=2)
    offline_items = [
        'When Wi-Fi or Mobile Data is disconnected, the app displays cached weather and telemetry.',
        'Mood logs submitted while offline are saved to the local Hive sync queue.',
        'When connectivity is restored, the app automatically flushes the queue to Cloud Firestore.',
    ]
    for f in offline_items:
        doc.add_paragraph(f, style='List Bullet')

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 11: SUMMARY
    # =========================================================================
    doc.add_heading('11. Summary & User Support', level=1)
    doc.add_paragraph(
        'MindSync AI combines state-of-the-art machine learning, AI coaching, and offline security '
        'to deliver a complete mental wellness companion for IT professionals.'
    )

    add_styled_table(doc,
        ['Need Help?', 'Contact / Path'],
        [
            ['Support Email', 'support@mindsyncai.com'],
            ['API Documentation', 'http://localhost:8000/docs (FastAPI Swagger)'],
            ['Source Code & Docs', 'docs/ directory in repository'],
        ]
    )

    doc.add_paragraph('\n\n')
    p = doc.add_paragraph()
    run = p.add_run('© 2026 MindSync AI – AI Mental Wellness Companion for IT Professionals')
    run.italic = True
    run.font.color.rgb = RGBColor(150, 150, 150)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save
    output_path = os.path.join(os.path.dirname(__file__), 'MindSync_AI_User_Guide_With_Screenshots.docx')
    doc.save(output_path)
    print(f"\n[OK] User Guide saved successfully to: {output_path}")
    return output_path

if __name__ == '__main__':
    create_user_guide()
